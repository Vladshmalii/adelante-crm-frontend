"""Склад: товары, категории, движения остатков, импорт/экспорт, расходники записи.

Права (решение от 28.09.2026, web/docs/BACKEND.md):
- раздел целиком, включая цены товаров, — администратор и суперюзер;
- мастер — только выбор товара при списании (название, единица, остаток) и
  списание расходников по своей записи.

Остаток меняется только движениями; под блокировкой строки товара (FOR
UPDATE), чтобы параллельные списания не увели остаток в минус.
"""

import io
import uuid
from datetime import datetime
from decimal import Decimal, InvalidOperation
from enum import StrEnum
from typing import Annotated, Any

from fastapi import APIRouter, Depends, HTTPException, Query, UploadFile, status
from fastapi.responses import StreamingResponse
from pydantic import Field
from sqlalchemy import ColumnElement, case, func, or_, select, update
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.admin.deps import CurrentAuthor, ensure_own_record
from app.api.schemas import ApiModel, Envelope, PersonRef, page_meta
from app.api.security import AdminUser, CurrentUser, require_salon_access
from app.models.shard import (
    AuditAction,
    InventoryCategory,
    MovementType,
    Product,
    ProductUnit,
    Record,
    RecordStatus,
    StockMovement,
)
from app.services.audit import diff_fields, write_audit
from app.tenancy.deps import TenantSession

router = APIRouter(tags=["inventory"], dependencies=[Depends(require_salon_access)])

UNIT_NAMES = {
    ProductUnit.PCS: "шт",
    ProductUnit.ML: "мл",
    ProductUnit.L: "л",
    ProductUnit.G: "г",
    ProductUnit.KG: "кг",
}


class StockStatus(StrEnum):
    IN_STOCK = "in_stock"
    LOW = "low"
    OUT = "out"


def _stock_status(product: Product) -> StockStatus:
    if product.quantity <= 0:
        return StockStatus.OUT
    if product.quantity <= product.min_quantity:
        return StockStatus.LOW
    return StockStatus.IN_STOCK


def _stock_status_filter(value: StockStatus) -> ColumnElement[bool]:
    if value == StockStatus.OUT:
        return Product.quantity <= 0
    if value == StockStatus.LOW:
        return (Product.quantity > 0) & (Product.quantity <= Product.min_quantity)
    return Product.quantity > Product.min_quantity


# --- Схемы ------------------------------------------------------------------


class ProductOut(ApiModel):
    id: uuid.UUID
    name: str
    unit: ProductUnit
    quantity: Decimal
    stock_status: StockStatus
    # Для мастера (выбор при списании) — null
    sku: str | None = None
    category: PersonRef | None = None
    package_volume: Decimal | None = None
    min_quantity: Decimal | None = None
    cost_price: Decimal | None = None
    sale_price: Decimal | None = None
    description: str | None = None
    is_active: bool = True
    created_at: datetime | None = None


def _product_out(product: Product, categories: dict[uuid.UUID, str], *, full: bool) -> ProductOut:
    out = ProductOut(
        id=product.id,
        name=product.name,
        unit=product.unit,
        quantity=product.quantity,
        stock_status=_stock_status(product),
    )
    if full:
        out.sku = product.sku
        out.category = PersonRef(
            id=str(product.category_id), name=categories.get(product.category_id)
        )
        out.package_volume = product.package_volume
        out.min_quantity = product.min_quantity
        out.cost_price = product.cost_price
        out.sale_price = product.sale_price
        out.description = product.description
        out.is_active = product.is_active
        out.created_at = product.created_at
    return out


async def _category_names(tenant_session: AsyncSession) -> dict[uuid.UUID, str]:
    rows = await tenant_session.execute(select(InventoryCategory.id, InventoryCategory.name))
    return {category_id: name for category_id, name in rows.all()}


async def _system_category(tenant_session: AsyncSession) -> InventoryCategory:
    category = await tenant_session.scalar(
        select(InventoryCategory).where(InventoryCategory.is_system.is_(True))
    )
    if category is None:
        category = InventoryCategory(name="Без категорії", is_system=True)
        tenant_session.add(category)
        await tenant_session.flush()
    return category


async def _get_product(
    tenant_session: AsyncSession, product_id: uuid.UUID, *, for_update: bool = False
) -> Product:
    query = select(Product).where(Product.id == product_id)
    if for_update:
        query = query.with_for_update()
    product = await tenant_session.scalar(query)
    if product is None:
        raise HTTPException(404, "Товар не найден")
    return product


def _product_title(product: Product) -> str:
    return f"Товар «{product.name}» ({product.sku})"


# --- Товары -----------------------------------------------------------------


SORTS = {
    "name": Product.name,
    "sku": Product.sku,
    "quantity": Product.quantity,
    "createdAt": Product.created_at,
}


@router.get("/inventory/products", response_model=Envelope[list[ProductOut]])
async def list_products(
    user: CurrentUser,
    tenant_session: TenantSession,
    query_text: Annotated[str | None, Query(alias="query")] = None,
    category_id: Annotated[uuid.UUID | None, Query(alias="categoryId")] = None,
    stock_status: Annotated[StockStatus | None, Query(alias="stockStatus")] = None,
    include_inactive: Annotated[bool, Query(alias="includeInactive")] = False,
    sort: str = "name",
    desc: bool = False,
    page: int = 1,
    per_page: Annotated[int, Query(alias="perPage", le=500)] = 50,
) -> Envelope[list[ProductOut]]:
    """Список товаров. Мастеру — только активные и только название, единица, остаток."""
    query = select(Product)
    if user.is_master or not include_inactive:
        query = query.where(Product.is_active.is_(True))
    if query_text:
        pattern = f"%{query_text}%"
        query = query.where(or_(Product.name.ilike(pattern), Product.sku.ilike(pattern)))
    if category_id is not None:
        query = query.where(Product.category_id == category_id)
    if stock_status is not None:
        query = query.where(_stock_status_filter(stock_status))

    total = await tenant_session.scalar(select(func.count()).select_from(query.subquery()))
    column = SORTS.get(sort, Product.name)
    products = await tenant_session.scalars(
        query.order_by(column.desc() if desc else column, Product.name)
        .offset((page - 1) * per_page)
        .limit(per_page)
    )
    categories = await _category_names(tenant_session) if user.is_admin else {}
    return Envelope(
        data=[_product_out(p, categories, full=user.is_admin) for p in products],
        meta=page_meta(page, per_page, total or 0),
    )


class InventorySummaryOut(ApiModel):
    total: int
    low: int
    out: int
    stock_value: Decimal


@router.get("/inventory/summary", response_model=Envelope[InventorySummaryOut])
async def inventory_summary(
    _admin: AdminUser, tenant_session: TenantSession
) -> Envelope[InventorySummaryOut]:
    total, low, out, value = (
        await tenant_session.execute(
            select(
                func.count(),
                func.count().filter(
                    (Product.quantity > 0) & (Product.quantity <= Product.min_quantity)
                ),
                func.count().filter(Product.quantity <= 0),
                func.coalesce(
                    func.sum(
                        case(
                            (Product.quantity > 0, Product.quantity * Product.cost_price),
                            else_=0,
                        )
                    ),
                    0,
                ),
            ).where(Product.is_active.is_(True))
        )
    ).one()
    return Envelope(
        data=InventorySummaryOut(
            total=total, low=low, out=out, stock_value=Decimal(value).quantize(Decimal("0.01"))
        )
    )


@router.get("/inventory/products/{product_id}", response_model=Envelope[ProductOut])
async def get_product(
    product_id: uuid.UUID, _admin: AdminUser, tenant_session: TenantSession
) -> Envelope[ProductOut]:
    product = await _get_product(tenant_session, product_id)
    return Envelope(data=_product_out(product, await _category_names(tenant_session), full=True))


class ProductCreateIn(ApiModel):
    name: str = Field(min_length=1, max_length=255)
    sku: str = Field(min_length=1, max_length=64)
    # Не указана — «Без категорії»
    category_id: uuid.UUID | None = None
    unit: ProductUnit = ProductUnit.PCS
    package_volume: Decimal | None = Field(default=None, gt=0)
    # Начальный остаток — оформляется движением «надходження»
    quantity: Decimal = Field(default=Decimal(0), ge=0)
    min_quantity: Decimal = Field(default=Decimal(0), ge=0)
    cost_price: Decimal | None = Field(default=None, ge=0)
    sale_price: Decimal | None = Field(default=None, ge=0)
    description: str | None = Field(default=None, max_length=2000)


async def _ensure_sku_free(
    tenant_session: AsyncSession, sku: str, exclude_id: uuid.UUID | None = None
) -> None:
    query = select(Product.id).where(func.lower(Product.sku) == sku.lower())
    if exclude_id is not None:
        query = query.where(Product.id != exclude_id)
    if await tenant_session.scalar(query) is not None:
        raise HTTPException(409, "Товар с таким артикулом уже есть")


async def _ensure_category(tenant_session: AsyncSession, category_id: uuid.UUID) -> None:
    if await tenant_session.get(InventoryCategory, category_id) is None:
        raise HTTPException(422, "Категория не найдена")


def _add_movement(
    tenant_session: AsyncSession,
    product: Product,
    *,
    type_: MovementType,
    delta: Decimal,
    reason: str | None,
    author_id: uuid.UUID | None,
    author_name: str | None,
    record_id: uuid.UUID | None = None,
) -> StockMovement:
    product.quantity = product.quantity + delta
    movement = StockMovement(
        product_id=product.id,
        type=type_,
        delta=delta,
        quantity_after=product.quantity,
        reason=reason,
        record_id=record_id,
        author_id=author_id,
        author_name=author_name,
    )
    tenant_session.add(movement)
    return movement


@router.post(
    "/inventory/products", response_model=Envelope[ProductOut], status_code=status.HTTP_201_CREATED
)
async def create_product(
    body: ProductCreateIn,
    _admin: AdminUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
) -> Envelope[ProductOut]:
    await _ensure_sku_free(tenant_session, body.sku)
    if body.category_id is not None:
        await _ensure_category(tenant_session, body.category_id)
        category_id = body.category_id
    else:
        category_id = (await _system_category(tenant_session)).id

    product = Product(
        **body.model_dump(by_alias=False, exclude={"quantity", "category_id"}),
        category_id=category_id,
        quantity=Decimal(0),
    )
    tenant_session.add(product)
    await tenant_session.flush()
    if body.quantity > 0:
        _add_movement(
            tenant_session,
            product,
            type_=MovementType.RECEIPT,
            delta=body.quantity,
            reason="Початковий залишок",
            author_id=author.id,
            author_name=author.name,
        )
    write_audit(
        tenant_session,
        entity="product",
        entity_id=product.id,
        entity_name=_product_title(product),
        action=AuditAction.CREATED,
        author_id=author.id,
        author_name=author.name,
    )
    await tenant_session.flush()
    await tenant_session.refresh(product, attribute_names=["created_at"])
    return Envelope(data=_product_out(product, await _category_names(tenant_session), full=True))


class ProductPatchIn(ApiModel):
    """Остаток здесь не меняется — только движениями."""

    name: str | None = Field(default=None, min_length=1, max_length=255)
    sku: str | None = Field(default=None, min_length=1, max_length=64)
    category_id: uuid.UUID | None = None
    unit: ProductUnit | None = None
    package_volume: Decimal | None = Field(default=None, gt=0)
    min_quantity: Decimal | None = Field(default=None, ge=0)
    cost_price: Decimal | None = Field(default=None, ge=0)
    sale_price: Decimal | None = Field(default=None, ge=0)
    description: str | None = Field(default=None, max_length=2000)
    is_active: bool | None = None


@router.patch("/inventory/products/{product_id}", response_model=Envelope[ProductOut])
async def patch_product(
    product_id: uuid.UUID,
    body: ProductPatchIn,
    _admin: AdminUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
) -> Envelope[ProductOut]:
    product = await _get_product(tenant_session, product_id)
    updates = body.model_dump(exclude_unset=True, by_alias=False)
    for required in ("name", "sku", "category_id", "unit", "min_quantity", "is_active"):
        if required in updates and updates[required] is None:
            raise HTTPException(422, f"Поле {required} нельзя очистить")
    if updates.get("sku"):
        await _ensure_sku_free(tenant_session, updates["sku"], exclude_id=product.id)
    if updates.get("category_id"):
        await _ensure_category(tenant_session, updates["category_id"])

    changes = diff_fields(product, updates)
    for field, value in updates.items():
        setattr(product, field, value)
    if changes:
        write_audit(
            tenant_session,
            entity="product",
            entity_id=product.id,
            entity_name=_product_title(product),
            action=AuditAction.UPDATED,
            author_id=author.id,
            author_name=author.name,
            details=changes,
        )
    return Envelope(data=_product_out(product, await _category_names(tenant_session), full=True))


@router.delete("/inventory/products/{product_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_product(
    product_id: uuid.UUID,
    _admin: AdminUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
) -> None:
    """Мягкое удаление: на товар ссылается история движений."""
    product = await _get_product(tenant_session, product_id)
    product.is_active = False
    write_audit(
        tenant_session,
        entity="product",
        entity_id=product.id,
        entity_name=_product_title(product),
        action=AuditAction.DELETED,
        author_id=author.id,
        author_name=author.name,
    )


# --- Движения ---------------------------------------------------------------


class MovementIn(ApiModel):
    type: MovementType
    # receipt / write_off — количество (> 0); adjustment — фактический остаток (≥ 0)
    quantity: Decimal = Field(ge=0)
    reason: str | None = Field(default=None, max_length=500)


class MovementOut(ApiModel):
    id: uuid.UUID
    product_id: uuid.UUID
    type: MovementType
    delta: Decimal
    quantity_after: Decimal
    reason: str | None
    record_id: uuid.UUID | None
    author: PersonRef
    created_at: datetime


def _movement_out(m: StockMovement) -> MovementOut:
    return MovementOut(
        id=m.id,
        product_id=m.product_id,
        type=m.type,
        delta=m.delta,
        quantity_after=m.quantity_after,
        reason=m.reason,
        record_id=m.record_id,
        author=PersonRef(id=str(m.author_id) if m.author_id else None, name=m.author_name),
        created_at=m.created_at,
    )


def _movement_delta(product: Product, body: MovementIn) -> Decimal:
    if body.type == MovementType.ADJUSTMENT:
        return body.quantity - product.quantity
    if body.quantity <= 0:
        raise HTTPException(422, "Количество должно быть больше нуля")
    if body.type == MovementType.WRITE_OFF:
        if body.quantity > product.quantity:
            raise HTTPException(
                409, f"Недостаточно товара: в наличии {product.quantity} {UNIT_NAMES[product.unit]}"
            )
        return -body.quantity
    return body.quantity


async def _saved_movement(tenant_session: AsyncSession, movement: StockMovement) -> MovementOut:
    await tenant_session.flush()
    await tenant_session.refresh(movement, attribute_names=["created_at"])
    return _movement_out(movement)


@router.post(
    "/inventory/products/{product_id}/movements",
    response_model=Envelope[MovementOut],
    status_code=status.HTTP_201_CREATED,
)
async def create_movement(
    product_id: uuid.UUID,
    body: MovementIn,
    _admin: AdminUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
) -> Envelope[MovementOut]:
    product = await _get_product(tenant_session, product_id, for_update=True)
    if not product.is_active:
        raise HTTPException(409, "Товар удалён")
    delta = _movement_delta(product, body)
    movement = _add_movement(
        tenant_session,
        product,
        type_=body.type,
        delta=delta,
        reason=body.reason,
        author_id=author.id,
        author_name=author.name,
    )
    write_audit(
        tenant_session,
        entity="product",
        entity_id=product.id,
        entity_name=_product_title(product),
        action=AuditAction.UPDATED,
        author_id=author.id,
        author_name=author.name,
        details={"quantity": [str(product.quantity - delta), str(product.quantity)]},
    )
    return Envelope(data=await _saved_movement(tenant_session, movement))


@router.get(
    "/inventory/products/{product_id}/movements", response_model=Envelope[list[MovementOut]]
)
async def list_movements(
    product_id: uuid.UUID,
    _admin: AdminUser,
    tenant_session: TenantSession,
    page: int = 1,
    per_page: Annotated[int, Query(alias="perPage", le=200)] = 50,
) -> Envelope[list[MovementOut]]:
    await _get_product(tenant_session, product_id)
    query = select(StockMovement).where(StockMovement.product_id == product_id)
    total = await tenant_session.scalar(select(func.count()).select_from(query.subquery()))
    movements = await tenant_session.scalars(
        query.order_by(StockMovement.created_at.desc())
        .offset((page - 1) * per_page)
        .limit(per_page)
    )
    return Envelope(
        data=[_movement_out(m) for m in movements], meta=page_meta(page, per_page, total or 0)
    )


# --- Категории --------------------------------------------------------------


class CategoryOut(ApiModel):
    id: uuid.UUID
    name: str
    is_system: bool
    products_count: int


@router.get("/inventory/categories", response_model=Envelope[list[CategoryOut]])
async def list_categories(
    _admin: AdminUser, tenant_session: TenantSession
) -> Envelope[list[CategoryOut]]:
    counts = dict(
        (
            await tenant_session.execute(
                select(Product.category_id, func.count())
                .where(Product.is_active.is_(True))
                .group_by(Product.category_id)
            )
        ).all()
    )
    categories = await tenant_session.scalars(
        select(InventoryCategory).order_by(
            InventoryCategory.is_system.desc(), InventoryCategory.name
        )
    )
    return Envelope(
        data=[
            CategoryOut(
                id=c.id, name=c.name, is_system=c.is_system, products_count=counts.get(c.id, 0)
            )
            for c in categories
        ]
    )


class CategoryIn(ApiModel):
    name: str = Field(min_length=1, max_length=128)


async def _ensure_category_name_free(
    tenant_session: AsyncSession, name: str, exclude_id: uuid.UUID | None = None
) -> None:
    query = select(InventoryCategory.id).where(
        func.lower(InventoryCategory.name) == name.strip().lower()
    )
    if exclude_id is not None:
        query = query.where(InventoryCategory.id != exclude_id)
    if await tenant_session.scalar(query) is not None:
        raise HTTPException(409, "Категория с таким названием уже есть")


@router.post(
    "/inventory/categories",
    response_model=Envelope[CategoryOut],
    status_code=status.HTTP_201_CREATED,
)
async def create_category(
    body: CategoryIn,
    _admin: AdminUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
) -> Envelope[CategoryOut]:
    await _ensure_category_name_free(tenant_session, body.name)
    category = InventoryCategory(name=body.name.strip(), is_system=False)
    tenant_session.add(category)
    await tenant_session.flush()
    write_audit(
        tenant_session,
        entity="product",
        entity_id=category.id,
        entity_name=f"Категорія складу «{category.name}»",
        action=AuditAction.CREATED,
        author_id=author.id,
        author_name=author.name,
    )
    return Envelope(
        data=CategoryOut(id=category.id, name=category.name, is_system=False, products_count=0)
    )


async def _get_editable_category(
    tenant_session: AsyncSession, category_id: uuid.UUID
) -> InventoryCategory:
    category = await tenant_session.get(InventoryCategory, category_id)
    if category is None:
        raise HTTPException(404, "Категория не найдена")
    if category.is_system:
        raise HTTPException(409, "Системную категорию нельзя изменить или удалить")
    return category


@router.patch("/inventory/categories/{category_id}", response_model=Envelope[CategoryOut])
async def rename_category(
    category_id: uuid.UUID,
    body: CategoryIn,
    _admin: AdminUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
) -> Envelope[CategoryOut]:
    category = await _get_editable_category(tenant_session, category_id)
    await _ensure_category_name_free(tenant_session, body.name, exclude_id=category.id)
    old = category.name
    category.name = body.name.strip()
    write_audit(
        tenant_session,
        entity="product",
        entity_id=category.id,
        entity_name=f"Категорія складу «{category.name}»",
        action=AuditAction.UPDATED,
        author_id=author.id,
        author_name=author.name,
        details={"name": [old, category.name]},
    )
    count = await tenant_session.scalar(
        select(func.count()).where(Product.category_id == category.id, Product.is_active.is_(True))
    )
    return Envelope(
        data=CategoryOut(
            id=category.id, name=category.name, is_system=False, products_count=count or 0
        )
    )


class CategoryDeleteMode(StrEnum):
    DELETE_PRODUCTS = "delete_products"
    MOVE_TO_UNCATEGORIZED = "move_to_uncategorized"


@router.delete("/inventory/categories/{category_id}", status_code=status.HTTP_204_NO_CONTENT)
async def delete_category(
    category_id: uuid.UUID,
    _admin: AdminUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
    mode: CategoryDeleteMode = CategoryDeleteMode.MOVE_TO_UNCATEGORIZED,
) -> None:
    """Удаление категории. Товары переносятся в «Без категорії»; при
    mode=delete_products они ещё и удаляются (мягко — история движений остаётся).
    """
    category = await _get_editable_category(tenant_session, category_id)
    uncategorized = await _system_category(tenant_session)
    values: dict[str, Any] = {"category_id": uncategorized.id}
    if mode == CategoryDeleteMode.DELETE_PRODUCTS:
        values["is_active"] = False
    moved = await tenant_session.execute(
        update(Product).where(Product.category_id == category.id).values(**values)
    )
    await tenant_session.delete(category)
    write_audit(
        tenant_session,
        entity="product",
        entity_id=category.id,
        entity_name=f"Категорія складу «{category.name}»",
        action=AuditAction.DELETED,
        author_id=author.id,
        author_name=author.name,
        details={"products": [moved.rowcount, mode.value]},  # type: ignore[attr-defined]
    )


# --- Импорт / экспорт -------------------------------------------------------

IMPORT_COLUMNS = [
    "Назва",
    "Артикул",
    "Категорія",
    "Одиниця",
    "Об'єм упаковки",
    "Залишок",
    "Мін. залишок",
    "Собівартість",
    "Ціна продажу",
    "Опис",
]
UNITS_BY_NAME = {
    **{name: unit for unit, name in UNIT_NAMES.items()},
    **{u.value: u for u in ProductUnit},
}


class ImportReportOut(ApiModel):
    created: int
    updated: int
    errors: list[str]


def _decimal(value: Any, field: str, row: int) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        result = Decimal(str(value).replace(",", ".").replace(" ", ""))
    except InvalidOperation:
        raise ValueError(f"Рядок {row}: «{field}» — не число") from None
    if result < 0:
        raise ValueError(f"Рядок {row}: «{field}» не може бути від'ємним")
    return result


@router.post("/inventory/import", response_model=Envelope[ImportReportOut])
async def import_products(
    file: UploadFile,
    _admin: AdminUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
) -> Envelope[ImportReportOut]:
    """Импорт из Excel (колонки — IMPORT_COLUMNS, первая строка — заголовок).

    Товар ищется по артикулу: новый — создаётся, существующий — обновляется;
    отличие остатка оформляется коригуванням. Неизвестная категория создаётся.
    """
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(await file.read()), read_only=True, data_only=True)
    except Exception:  # noqa: BLE001 - openpyxl raises assorted exceptions for malformed uploads
        raise HTTPException(422, "Не удалось прочитать файл — ожидается .xlsx")

    categories = {
        c.name.lower(): c for c in await tenant_session.scalars(select(InventoryCategory))
    }
    uncategorized = await _system_category(tenant_session)
    created = updated = 0
    errors: list[str] = []
    for i, raw in enumerate(wb.active.iter_rows(min_row=2, values_only=True), start=2):
        row = (list(raw) + [None] * len(IMPORT_COLUMNS))[: len(IMPORT_COLUMNS)]
        if all(v is None or str(v).strip() == "" for v in row):
            continue
        name, sku, category_name, unit_name = (
            str(v).strip() if v is not None else "" for v in row[:4]
        )
        if not name or not sku:
            errors.append(f"Рядок {i}: потрібні назва і артикул")
            continue
        unit = UNITS_BY_NAME.get(unit_name.lower() or "шт")
        if unit is None:
            errors.append(f"Рядок {i}: невідома одиниця «{unit_name}»")
            continue
        try:
            volume, quantity, min_quantity, cost, price = (
                _decimal(row[4 + k], IMPORT_COLUMNS[4 + k], i) for k in range(5)
            )
        except ValueError as exc:
            errors.append(str(exc))
            continue

        category = uncategorized
        if category_name:
            found = categories.get(category_name.lower())
            if found is None:
                found = InventoryCategory(name=category_name, is_system=False)
                tenant_session.add(found)
                await tenant_session.flush()
                categories[category_name.lower()] = found
            category = found

        fields = {
            "name": name,
            "category_id": category.id,
            "unit": unit,
            "package_volume": volume,
            "min_quantity": min_quantity or Decimal(0),
            "cost_price": cost,
            "sale_price": price,
            "description": str(row[9]).strip() if row[9] else None,
        }
        product = await tenant_session.scalar(
            select(Product).where(func.lower(Product.sku) == sku.lower()).with_for_update()
        )
        if product is None:
            product = Product(sku=sku, quantity=Decimal(0), is_active=True, **fields)
            tenant_session.add(product)
            await tenant_session.flush()
            created += 1
        else:
            for field, value in fields.items():
                setattr(product, field, value)
            product.is_active = True
            updated += 1
        if quantity is not None and quantity != product.quantity:
            _add_movement(
                tenant_session,
                product,
                type_=(MovementType.RECEIPT if product.quantity == 0 else MovementType.ADJUSTMENT),
                delta=quantity - product.quantity,
                reason="Імпорт з Excel",
                author_id=author.id,
                author_name=author.name,
            )

    write_audit(
        tenant_session,
        entity="product",
        entity_id="import",
        entity_name=f"Імпорт складу з Excel: створено {created}, оновлено {updated}",
        action=AuditAction.CREATED,
        author_id=author.id,
        author_name=author.name,
        details={"created": [None, created], "updated": [None, updated]},
    )
    return Envelope(data=ImportReportOut(created=created, updated=updated, errors=errors))


EXPORT_BLOCKS = {
    "main": ["Назва", "Артикул", "Категорія", "Одиниця"],
    "stock": ["Об'єм упаковки", "Залишок", "Мін. залишок", "Статус"],
    "finance": ["Собівартість", "Ціна продажу", "Вартість залишку"],
    "description": ["Опис"],
}
STATUS_NAMES = {
    StockStatus.IN_STOCK: "В наявності",
    StockStatus.LOW: "Закінчується",
    StockStatus.OUT: "Немає",
}


def _export_values(block: str, p: Product, categories: dict[uuid.UUID, str]) -> list[Any]:
    if block == "main":
        return [p.name, p.sku, categories.get(p.category_id), UNIT_NAMES[p.unit]]
    if block == "stock":
        return [
            float(p.package_volume) if p.package_volume is not None else None,
            float(p.quantity),
            float(p.min_quantity),
            STATUS_NAMES[_stock_status(p)],
        ]
    if block == "finance":
        return [
            float(p.cost_price) if p.cost_price is not None else None,
            float(p.sale_price) if p.sale_price is not None else None,
            float(p.cost_price * p.quantity) if p.cost_price is not None else None,
        ]
    return [p.description]


@router.get("/inventory/export")
async def export_products(
    _admin: AdminUser,
    tenant_session: TenantSession,
    blocks: str = "main,stock,finance,description",
    category_id: Annotated[uuid.UUID | None, Query(alias="categoryId")] = None,
) -> StreamingResponse:
    """Excel-выгрузка склада; blocks — через запятую из main, stock, finance, description."""
    from openpyxl import Workbook

    chosen = ["main"] + [b for b in EXPORT_BLOCKS if b != "main" and b in blocks.split(",")]
    query = select(Product).where(Product.is_active.is_(True)).order_by(Product.name)
    if category_id is not None:
        query = query.where(Product.category_id == category_id)
    products = list(await tenant_session.scalars(query))
    categories = await _category_names(tenant_session)

    wb = Workbook()
    ws = wb.active
    ws.title = "Склад"
    ws.append([title for block in chosen for title in EXPORT_BLOCKS[block]])
    for p in products:
        ws.append([v for block in chosen for v in _export_values(block, p, categories)])

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="inventory.xlsx"'},
    )


# --- Расходники записи ------------------------------------------------------


class ConsumableOut(ApiModel):
    movement_id: uuid.UUID
    product_id: uuid.UUID
    product_name: str
    unit: ProductUnit
    quantity: Decimal
    author: PersonRef
    created_at: datetime


class ConsumableIn(ApiModel):
    product_id: uuid.UUID
    quantity: Decimal = Field(gt=0)


class ConsumablesIn(ApiModel):
    items: list[ConsumableIn] = Field(min_length=1)


async def _record_for_consumables(
    tenant_session: AsyncSession, record_id: uuid.UUID, user: CurrentUser
) -> Record:
    record = await tenant_session.get(Record, record_id)
    if record is None:
        raise HTTPException(404, "Запись не найдена")
    ensure_own_record(user, record)
    return record


async def _consumables(tenant_session: AsyncSession, record_id: uuid.UUID) -> list[ConsumableOut]:
    rows = await tenant_session.execute(
        select(StockMovement, Product)
        .join(Product, Product.id == StockMovement.product_id)
        .where(StockMovement.record_id == record_id)
        .order_by(StockMovement.created_at)
    )
    return [
        ConsumableOut(
            movement_id=m.id,
            product_id=p.id,
            product_name=p.name,
            unit=p.unit,
            quantity=-m.delta,
            author=PersonRef(id=str(m.author_id) if m.author_id else None, name=m.author_name),
            created_at=m.created_at,
        )
        for m, p in rows.all()
    ]


@router.get("/records/{record_id}/consumables", response_model=Envelope[list[ConsumableOut]])
async def list_consumables(
    record_id: uuid.UUID, user: CurrentUser, tenant_session: TenantSession
) -> Envelope[list[ConsumableOut]]:
    await _record_for_consumables(tenant_session, record_id, user)
    return Envelope(data=await _consumables(tenant_session, record_id))


@router.post(
    "/records/{record_id}/consumables",
    response_model=Envelope[list[ConsumableOut]],
    status_code=status.HTTP_201_CREATED,
)
async def write_off_consumables(
    record_id: uuid.UUID,
    body: ConsumablesIn,
    user: CurrentUser,
    author: CurrentAuthor,
    tenant_session: TenantSession,
) -> Envelope[list[ConsumableOut]]:
    """Списание расходников по записи (мастер — по своей записи, администратор — по любой)."""
    record = await _record_for_consumables(tenant_session, record_id, user)
    if record.status == RecordStatus.CANCELLED:
        raise HTTPException(409, "По отменённой записи списывать нельзя")

    # Одинаковые товары суммируются; блокировки — в порядке id (без взаимоблокировок)
    totals: dict[uuid.UUID, Decimal] = {}
    for item in body.items:
        totals[item.product_id] = totals.get(item.product_id, Decimal(0)) + item.quantity
    for product_id in sorted(totals):
        product = await _get_product(tenant_session, product_id, for_update=True)
        if not product.is_active:
            raise HTTPException(409, f"Товар «{product.name}» удалён")
        if totals[product_id] > product.quantity:
            raise HTTPException(
                409,
                f"Недостаточно «{product.name}»: в наличии {product.quantity} "
                f"{UNIT_NAMES[product.unit]}",
            )
        _add_movement(
            tenant_session,
            product,
            type_=MovementType.WRITE_OFF,
            delta=-totals[product_id],
            reason=f"Витрати на візит: {record.client_name}",
            author_id=author.id,
            author_name=author.name,
            record_id=record.id,
        )
    write_audit(
        tenant_session,
        entity="record",
        entity_id=record.id,
        entity_name=f"{record.client_name} → {record.master_name or 'без майстра'}",
        action=AuditAction.UPDATED,
        author_id=author.id,
        author_name=author.name,
        details={"consumables": [None, len(totals)]},
    )
    await tenant_session.flush()
    return Envelope(data=await _consumables(tenant_session, record.id))
