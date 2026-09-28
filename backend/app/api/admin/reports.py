"""Звіти: сводка, выручка, клиенты, сотрудники, услуги, Excel.

Выручка = сумма неотменённых чеков по дате чека. Выручка мастера и услуги —
по чекам, привязанным к записям; выручка записи делится между её услугами
пропорционально их цене (снапшот в record_services).

Периоды и группировка — по киевским суткам (неделя с понедельника).

Права: раздел — администратор и суперюзер; денежные показатели — только
суперюзер (администратору — null, GET /reports/revenue — 403).
"""

import io
import uuid
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from enum import StrEnum
from typing import Annotated, Any

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import distinct, func, select
from sqlalchemy.ext.asyncio import AsyncSession

from app.api.schemas import ApiModel, Envelope
from app.api.security import AdminUser, SuperUser, require_admin
from app.models.shard import (
    Receipt,
    ReceiptStatus,
    Record,
    RecordService,
    RecordStatus,
    Review,
    Service,
)
from app.tenancy.deps import TenantSession
from app.timeutils import LocalDatetime, day_bounds, to_local

router = APIRouter(prefix="/reports", tags=["reports"], dependencies=[Depends(require_admin)])

CENT = Decimal("0.01")


class GroupBy(StrEnum):
    DAY = "day"
    WEEK = "week"
    MONTH = "month"


DateFrom = Annotated[LocalDatetime, Query(alias="dateFrom")]
DateTo = Annotated[LocalDatetime, Query(alias="dateTo")]


def _check_period(date_from: datetime, date_to: datetime) -> None:
    if date_to <= date_from:
        raise HTTPException(422, "dateTo має бути пізніше за dateFrom")
    if date_to - date_from > timedelta(days=731):
        raise HTTPException(422, "Період звіту — не більше двох років")


# --- Периоды ------------------------------------------------------------------


def _bucket_start(day: date, group: GroupBy) -> date:
    if group == GroupBy.WEEK:
        return day - timedelta(days=day.weekday())
    if group == GroupBy.MONTH:
        return day.replace(day=1)
    return day


def _next_bucket(start: date, group: GroupBy) -> date:
    if group == GroupBy.WEEK:
        return start + timedelta(days=7)
    if group == GroupBy.MONTH:
        return (start + timedelta(days=32)).replace(day=1)
    return start + timedelta(days=1)


def _previous_bucket(start: date, group: GroupBy) -> date:
    return _bucket_start(start - timedelta(days=1), group)


def _buckets(date_from: datetime, date_to: datetime, group: GroupBy) -> list[date]:
    """Начала интервалов, покрывающих [date_from, date_to) по киевскому времени."""
    last_day = to_local(date_to - timedelta(microseconds=1)).date()
    current = _bucket_start(to_local(date_from).date(), group)
    result = []
    while current <= last_day:
        result.append(current)
        current = _next_bucket(current, group)
    return result


def _bucket_of(value: datetime, group: GroupBy) -> date:
    return _bucket_start(to_local(value).date(), group)


def _change(current: Decimal | int, previous: Decimal | int) -> float | None:
    if not previous:
        return None
    return round(float((Decimal(current) - Decimal(previous)) / Decimal(previous) * 100), 1)


def _money(value: Decimal) -> Decimal:
    return Decimal(value).quantize(CENT)


# --- Данные -------------------------------------------------------------------


@dataclass
class _PaidReceipt:
    date: datetime
    amount: Decimal
    record_id: uuid.UUID | None


async def _paid_receipts(
    session: AsyncSession, date_from: datetime, date_to: datetime
) -> list[_PaidReceipt]:
    rows = await session.execute(
        select(Receipt.date, Receipt.amount, Receipt.record_id).where(
            Receipt.status != ReceiptStatus.CANCELLED,
            Receipt.date >= date_from,
            Receipt.date < date_to,
        )
    )
    return [_PaidReceipt(d, a, r) for d, a, r in rows.all()]


@dataclass
class _PeriodStats:
    revenue: Decimal
    receipts: int
    clients: int
    records: int

    @property
    def avg_check(self) -> Decimal:
        return _money(self.revenue / self.receipts) if self.receipts else Decimal(0)


async def _period_stats(
    session: AsyncSession, date_from: datetime, date_to: datetime
) -> _PeriodStats:
    revenue, receipts = (
        await session.execute(
            select(func.coalesce(func.sum(Receipt.amount), 0), func.count()).where(
                Receipt.status != ReceiptStatus.CANCELLED,
                Receipt.date >= date_from,
                Receipt.date < date_to,
            )
        )
    ).one()
    clients, records = (
        await session.execute(
            select(func.count(distinct(Record.client_id)), func.count()).where(
                Record.status != RecordStatus.CANCELLED,
                Record.start_at >= date_from,
                Record.start_at < date_to,
            )
        )
    ).one()
    return _PeriodStats(Decimal(revenue), receipts, clients, records)


# --- Сводка -------------------------------------------------------------------


class MetricOut(ApiModel):
    value: Decimal
    previous: Decimal
    change_percent: float | None


def _metric(current: Decimal | int, previous: Decimal | int) -> MetricOut:
    return MetricOut(
        value=Decimal(current),
        previous=Decimal(previous),
        change_percent=_change(current, previous),
    )


class SummaryOut(ApiModel):
    date_from: datetime
    date_to: datetime
    previous_from: datetime
    # Денежные — только суперюзеру
    revenue: MetricOut | None
    avg_check: MetricOut | None
    clients: MetricOut
    records: MetricOut


async def _summary(
    session: AsyncSession, date_from: datetime, date_to: datetime, *, show_money: bool
) -> SummaryOut:
    previous_from = date_from - (date_to - date_from)
    cur = await _period_stats(session, date_from, date_to)
    prev = await _period_stats(session, previous_from, date_from)
    return SummaryOut(
        date_from=date_from,
        date_to=date_to,
        previous_from=previous_from,
        revenue=_metric(_money(cur.revenue), _money(prev.revenue)) if show_money else None,
        avg_check=_metric(cur.avg_check, prev.avg_check) if show_money else None,
        clients=_metric(cur.clients, prev.clients),
        records=_metric(cur.records, prev.records),
    )


@router.get("/summary", response_model=Envelope[SummaryOut])
async def summary(
    user: AdminUser, tenant_session: TenantSession, date_from: DateFrom, date_to: DateTo
) -> Envelope[SummaryOut]:
    _check_period(date_from, date_to)
    return Envelope(
        data=await _summary(tenant_session, date_from, date_to, show_money=user.is_superuser)
    )


# --- Выручка ------------------------------------------------------------------


class RevenuePointOut(ApiModel):
    period: date
    amount: Decimal
    receipts: int


class RevenueOut(ApiModel):
    group_by: GroupBy
    total: Decimal
    points: list[RevenuePointOut]


async def _revenue(
    session: AsyncSession, date_from: datetime, date_to: datetime, group: GroupBy
) -> RevenueOut:
    amounts: dict[date, Decimal] = defaultdict(Decimal)
    counts: dict[date, int] = defaultdict(int)
    for receipt in await _paid_receipts(session, date_from, date_to):
        bucket = _bucket_of(receipt.date, group)
        amounts[bucket] += receipt.amount
        counts[bucket] += 1
    points = [
        RevenuePointOut(period=b, amount=_money(amounts[b]), receipts=counts[b])
        for b in _buckets(date_from, date_to, group)
    ]
    return RevenueOut(
        group_by=group, total=_money(sum(amounts.values(), Decimal(0))), points=points
    )


@router.get("/revenue", response_model=Envelope[RevenueOut])
async def revenue(
    _superuser: SuperUser,
    tenant_session: TenantSession,
    date_from: DateFrom,
    date_to: DateTo,
    group_by: Annotated[GroupBy, Query(alias="groupBy")] = GroupBy.DAY,
) -> Envelope[RevenueOut]:
    _check_period(date_from, date_to)
    return Envelope(data=await _revenue(tenant_session, date_from, date_to, group_by))


# --- Клиенты ------------------------------------------------------------------


class ClientsPointOut(ApiModel):
    period: date
    new: int
    returning: int
    total: int
    # % клиентов предыдущего интервала, у которых был визит в этом; null — не с чем сравнить
    retention_percent: float | None


class ClientsOut(ApiModel):
    group_by: GroupBy
    new: int
    returning: int
    total: int
    points: list[ClientsPointOut]


async def _clients(
    session: AsyncSession, date_from: datetime, date_to: datetime, group: GroupBy
) -> ClientsOut:
    """Новый — первый завершённый визит в салоне попал в интервал; остальные — повторные."""
    buckets = _buckets(date_from, date_to, group)
    lookback, _ = day_bounds(_previous_bucket(buckets[0], group))
    rows = (
        await session.execute(
            select(Record.client_id, Record.start_at).where(
                Record.status == RecordStatus.COMPLETED,
                Record.start_at >= min(lookback, date_from),
                Record.start_at < date_to,
            )
        )
    ).all()
    client_ids = {client_id for client_id, _ in rows}
    first_visit: dict[uuid.UUID, datetime] = {}
    if client_ids:
        first_visit = dict(
            (
                await session.execute(
                    select(Record.client_id, func.min(Record.start_at))
                    .where(
                        Record.status == RecordStatus.COMPLETED,
                        Record.client_id.in_(client_ids),
                    )
                    .group_by(Record.client_id)
                )
            ).all()
        )

    by_bucket: dict[date, set[uuid.UUID]] = defaultdict(set)
    in_period: set[uuid.UUID] = set()
    for client_id, start_at in rows:
        by_bucket[_bucket_of(start_at, group)].add(client_id)
        if date_from <= start_at < date_to:
            in_period.add(client_id)

    points = []
    for bucket in buckets:
        clients = by_bucket.get(bucket, set())
        new = {c for c in clients if _bucket_of(first_visit[c], group) == bucket}
        previous = by_bucket.get(_previous_bucket(bucket, group), set())
        points.append(
            ClientsPointOut(
                period=bucket,
                new=len(new),
                returning=len(clients) - len(new),
                total=len(clients),
                retention_percent=(
                    round(len(previous & clients) / len(previous) * 100, 1) if previous else None
                ),
            )
        )
    new_total = {c for c in in_period if date_from <= first_visit[c] < date_to}
    return ClientsOut(
        group_by=group,
        new=len(new_total),
        returning=len(in_period) - len(new_total),
        total=len(in_period),
        points=points,
    )


@router.get("/clients", response_model=Envelope[ClientsOut])
async def clients_report(
    _admin: AdminUser,
    tenant_session: TenantSession,
    date_from: DateFrom,
    date_to: DateTo,
    group_by: Annotated[GroupBy, Query(alias="groupBy")] = GroupBy.DAY,
) -> Envelope[ClientsOut]:
    _check_period(date_from, date_to)
    return Envelope(data=await _clients(tenant_session, date_from, date_to, group_by))


# --- Сотрудники ---------------------------------------------------------------


class StaffRowOut(ApiModel):
    master_id: uuid.UUID
    name: str
    records: int
    completed: int
    no_show: int
    cancelled: int
    # Денежные — только суперюзеру
    revenue: Decimal | None
    avg_check: Decimal | None
    rating: float | None
    reviews: int


async def _staff(
    session: AsyncSession, date_from: datetime, date_to: datetime, *, show_money: bool
) -> list[StaffRowOut]:
    counts = (
        await session.execute(
            select(
                Record.master_id,
                func.max(Record.master_name),
                func.count(),
                func.count().filter(Record.status == RecordStatus.COMPLETED),
                func.count().filter(Record.status == RecordStatus.NO_SHOW),
                func.count().filter(Record.status == RecordStatus.CANCELLED),
            )
            .where(
                Record.master_id.is_not(None),
                Record.start_at >= date_from,
                Record.start_at < date_to,
            )
            .group_by(Record.master_id)
        )
    ).all()
    money = {
        master_id: (amount, n)
        for master_id, amount, n in (
            await session.execute(
                select(Record.master_id, func.sum(Receipt.amount), func.count())
                .join(Record, Record.id == Receipt.record_id)
                .where(
                    Receipt.status != ReceiptStatus.CANCELLED,
                    Receipt.date >= date_from,
                    Receipt.date < date_to,
                    Record.master_id.is_not(None),
                )
                .group_by(Record.master_id)
            )
        ).all()
    }
    ratings = {
        master_id: (avg, n)
        for master_id, avg, n in (
            await session.execute(
                select(Review.master_id, func.avg(Review.rating), func.count())
                .where(Review.created_at >= date_from, Review.created_at < date_to)
                .group_by(Review.master_id)
            )
        ).all()
    }
    names: dict[uuid.UUID, str] = {}
    stats: dict[uuid.UUID, tuple[int, int, int, int]] = {}
    for master_id, name, total, completed, no_show, cancelled in counts:
        assert master_id is not None  # отфильтровано в запросе
        names[master_id] = name or "—"
        stats[master_id] = (total, completed, no_show, cancelled)
    # Мастер мог получить оплату в периоде за визит из прошлого периода
    for master_id in set(money) - set(names):
        assert master_id is not None
        names[master_id] = (
            await session.scalar(
                select(func.max(Record.master_name)).where(Record.master_id == master_id)
            )
        ) or "—"

    rows = []
    for master_id, name in names.items():
        total, completed, no_show, cancelled = stats.get(master_id, (0, 0, 0, 0))
        amount, paid = money.get(master_id, (Decimal(0), 0))
        avg_rating, reviews = ratings.get(master_id, (None, 0))
        rows.append(
            StaffRowOut(
                master_id=master_id,
                name=name,
                records=total,
                completed=completed,
                no_show=no_show,
                cancelled=cancelled,
                revenue=_money(amount) if show_money else None,
                avg_check=(_money(amount / paid) if paid else Decimal(0)) if show_money else None,
                rating=round(float(avg_rating), 2) if avg_rating is not None else None,
                reviews=reviews,
            )
        )
    rows.sort(key=lambda r: (-(r.revenue or 0), -r.completed, r.name))
    return rows


@router.get("/staff", response_model=Envelope[list[StaffRowOut]])
async def staff_report(
    user: AdminUser, tenant_session: TenantSession, date_from: DateFrom, date_to: DateTo
) -> Envelope[list[StaffRowOut]]:
    _check_period(date_from, date_to)
    return Envelope(
        data=await _staff(tenant_session, date_from, date_to, show_money=user.is_superuser)
    )


# --- Услуги -------------------------------------------------------------------


class ServiceRowOut(ApiModel):
    service_id: uuid.UUID
    name: str
    category: str
    count: int
    revenue: Decimal | None


class CategoryRowOut(ApiModel):
    category: str
    count: int
    revenue: Decimal | None


class ServicesOut(ApiModel):
    services: list[ServiceRowOut]
    categories: list[CategoryRowOut]


async def _services(
    session: AsyncSession, date_from: datetime, date_to: datetime, *, show_money: bool
) -> ServicesOut:
    counts = (
        await session.execute(
            select(RecordService.service_id, func.max(RecordService.name), func.count())
            .join(Record, Record.id == RecordService.record_id)
            .where(
                Record.status == RecordStatus.COMPLETED,
                Record.start_at >= date_from,
                Record.start_at < date_to,
            )
            .group_by(RecordService.service_id)
        )
    ).all()

    revenue: dict[uuid.UUID, Decimal] = defaultdict(Decimal)
    names: dict[uuid.UUID, str] = {service_id: name for service_id, name, _ in counts}
    paid = [r for r in await _paid_receipts(session, date_from, date_to) if r.record_id]
    if show_money and paid:
        items: dict[uuid.UUID, list[RecordService]] = defaultdict(list)
        for item in await session.scalars(
            select(RecordService).where(RecordService.record_id.in_({r.record_id for r in paid}))
        ):
            items[item.record_id].append(item)
        for receipt in paid:
            assert receipt.record_id is not None
            parts = items.get(receipt.record_id, [])
            base = sum((p.price for p in parts), Decimal(0))
            for part in parts:
                share = part.price / base if base else Decimal(1) / len(parts)
                revenue[part.service_id] += receipt.amount * share
                names.setdefault(part.service_id, part.name)

    service_ids = set(names)
    categories = (
        dict(
            (
                await session.execute(
                    select(Service.id, Service.category).where(Service.id.in_(service_ids))
                )
            ).all()
        )
        if service_ids
        else {}
    )
    count_by_id = {service_id: n for service_id, _, n in counts}
    services = [
        ServiceRowOut(
            service_id=service_id,
            name=names[service_id],
            category=categories.get(service_id, "other"),
            count=count_by_id.get(service_id, 0),
            revenue=_money(revenue[service_id]) if show_money else None,
        )
        for service_id in service_ids
    ]
    services.sort(key=lambda s: (-(s.revenue or 0), -s.count, s.name))

    by_category: dict[str, list[ServiceRowOut]] = defaultdict(list)
    for service in services:
        by_category[service.category].append(service)
    category_rows = [
        CategoryRowOut(
            category=category,
            count=sum(s.count for s in rows),
            revenue=_money(sum((s.revenue or Decimal(0) for s in rows), Decimal(0)))
            if show_money
            else None,
        )
        for category, rows in by_category.items()
    ]
    category_rows.sort(key=lambda c: (-(c.revenue or 0), -c.count, c.category))
    return ServicesOut(services=services, categories=category_rows)


@router.get("/services", response_model=Envelope[ServicesOut])
async def services_report(
    user: AdminUser, tenant_session: TenantSession, date_from: DateFrom, date_to: DateTo
) -> Envelope[ServicesOut]:
    _check_period(date_from, date_to)
    return Envelope(
        data=await _services(tenant_session, date_from, date_to, show_money=user.is_superuser)
    )


# --- Excel --------------------------------------------------------------------


def _num(value: Decimal | float | None) -> float | None:
    return float(value) if value is not None else None


@router.get("/export")
async def export_reports(
    user: AdminUser,
    tenant_session: TenantSession,
    date_from: DateFrom,
    date_to: DateTo,
    group_by: Annotated[GroupBy, Query(alias="groupBy")] = GroupBy.DAY,
) -> StreamingResponse:
    """Все отчёты одним файлом; администратору — без денежных колонок и листа выручки."""
    from openpyxl import Workbook

    _check_period(date_from, date_to)
    money = user.is_superuser
    wb = Workbook()

    ws = wb.active
    ws.title = "Підсумки"
    data = await _summary(tenant_session, date_from, date_to, show_money=money)
    ws.append(["Показник", "Період", "Попередній період", "Зміна, %"])
    metrics: list[tuple[str, MetricOut | None]] = [
        ("Виручка", data.revenue),
        ("Середній чек", data.avg_check),
        ("Клієнти", data.clients),
        ("Записи", data.records),
    ]
    for title, metric in metrics:
        if metric is not None:
            ws.append([title, _num(metric.value), _num(metric.previous), metric.change_percent])

    if money:
        ws = wb.create_sheet("Виручка")
        ws.append(["Період", "Виручка", "Чеків"])
        for point in (await _revenue(tenant_session, date_from, date_to, group_by)).points:
            ws.append([point.period.isoformat(), _num(point.amount), point.receipts])

    ws = wb.create_sheet("Клієнти")
    ws.append(["Період", "Нові", "Повторні", "Усього", "Утримання, %"])
    for c in (await _clients(tenant_session, date_from, date_to, group_by)).points:
        ws.append([c.period.isoformat(), c.new, c.returning, c.total, c.retention_percent])

    ws = wb.create_sheet("Співробітники")
    header: list[Any] = ["Майстер", "Записи", "Завершено", "Не прийшли", "Скасовано"]
    header += (["Виручка", "Середній чек"] if money else []) + ["Рейтинг", "Відгуків"]
    ws.append(header)
    for s in await _staff(tenant_session, date_from, date_to, show_money=money):
        row: list[Any] = [s.name, s.records, s.completed, s.no_show, s.cancelled]
        row += ([_num(s.revenue), _num(s.avg_check)] if money else []) + [s.rating, s.reviews]
        ws.append(row)

    ws = wb.create_sheet("Послуги")
    ws.append(["Послуга", "Категорія", "Кількість"] + (["Виручка"] if money else []))
    for svc in (await _services(tenant_session, date_from, date_to, show_money=money)).services:
        ws.append([svc.name, svc.category, svc.count] + ([_num(svc.revenue)] if money else []))

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return StreamingResponse(
        buffer,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": 'attachment; filename="reports.xlsx"'},
    )
