"""Звіти: суммы по чекам, сравнение с прошлым периодом, клиенты, мастера, услуги."""

import io
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from decimal import Decimal
from typing import Any
from urllib.parse import quote
from zoneinfo import ZoneInfo

import pytest
from openpyxl import load_workbook

from tests.conftest import Salon

TZ = ZoneInfo("Europe/Kyiv")


@dataclass
class World:
    s: Salon
    adm: str
    mst: str
    query: str


@pytest.fixture(scope="module")
def w(new_salon: Salon) -> World:
    s = new_salon
    api = s.api
    _, adm = s.create_staff("administrator")
    m1, mst = s.create_staff(firstName="Анна")
    m2, _ = s.create_staff(firstName="Богдана")
    cut = api.post(
        "/services",
        token=adm,
        salon=s.id,
        expect=201,
        json={
            "name": "Стрижка",
            "price": 300,
            "durationMinutes": 30,
            "categoryId": s.service_category("Волосся"),
        },
    )
    color = api.post(
        "/services",
        token=adm,
        salon=s.id,
        expect=201,
        json={
            "name": "Фарбування",
            "price": 900,
            "durationMinutes": 90,
            "categoryId": s.service_category("Фарбування"),
        },
    )
    reg = api.post(
        "/finances/cash-registers", token=s.su, salon=s.id, expect=201, json={"name": "Каса"}
    )
    pm = api.post(
        "/finances/payment-methods",
        token=s.su,
        salon=s.id,
        expect=201,
        json={"name": "Картка", "type": "card", "cashRegisterId": reg["id"]},
    )

    today = datetime.now(TZ).date()
    period_from = today - timedelta(days=6)

    def at(d: date, hour: int) -> str:
        return datetime(d.year, d.month, d.day, hour, tzinfo=TZ).isoformat()

    def make(
        master: dict[str, Any],
        services: list[dict[str, Any]],
        d: date,
        hour: int,
        client_phone: str,
        name: str,
        *,
        complete: bool = True,
        pay: bool = True,
    ) -> dict[str, Any]:
        rec = api.post(
            "/records",
            token=adm,
            salon=s.id,
            expect=201,
            json={
                "newClient": {"name": name, "phone": client_phone},
                "masterId": master["id"],
                "serviceIds": [x["id"] for x in services],
                "startAt": at(d, hour),
            },
        )
        if complete:
            api.post(f"/records/{rec['id']}/complete", token=adm, salon=s.id, json={})
        if complete and pay:
            api.post(
                f"/records/{rec['id']}/payment",
                token=adm,
                salon=s.id,
                json={
                    "payments": [{"paymentMethodId": pm["id"], "amount": float(rec["totalAmount"])}]
                },
            )
        return rec

    # Олена — первый визит до периода, в периоде повторная; её чек задним числом → прошлый период
    old = make(m1, [cut], today - timedelta(days=10), 10, "+380501119001", "Олена", pay=False)
    api.post(
        "/finances/receipts",
        token=s.su,
        salon=s.id,
        expect=201,
        json={
            "recordId": old["id"],
            "date": at(today - timedelta(days=10), 11),
            "payments": [{"paymentMethodId": pm["id"], "amount": 300}],
        },
    )
    a = make(m1, [cut, color], period_from + timedelta(days=1), 10, "+380501119001", "Олена")
    make(m2, [cut], period_from + timedelta(days=2), 11, "+380501119002", "Ірина")
    make(m1, [color], period_from + timedelta(days=3), 12, "+380501119003", "Софія", pay=False)
    d = make(
        m2, [cut], period_from + timedelta(days=3), 12, "+380501119004", "Дарина", complete=False
    )
    api.post(f"/records/{d['id']}/status", token=adm, salon=s.id, json={"status": "no_show"})

    # Отзыв о мастере 1 по токену завершённого визита
    from sqlalchemy import create_engine, text

    import cli

    with cli._master_session() as session:
        from app.models.master import Salon as SalonModel

        info = cli._conn_info(session.get(SalonModel, s.id))
    engine = create_engine(info.build_dsn(driver="psycopg"))
    with engine.connect() as conn:
        token = conn.execute(
            text("select review_token from records where id = :id"), {"id": a["id"]}
        ).scalar()
    engine.dispose()
    api.post(
        "/reviews",
        prefix=f"/api/booking/{s.slug}",
        expect=201,
        json={"token": str(token), "rating": 5},
    )

    start = datetime(period_from.year, period_from.month, period_from.day, tzinfo=TZ)
    end = start + timedelta(days=7)
    query = f"dateFrom={quote(start.isoformat())}&dateTo={quote(end.isoformat())}"
    return World(s, adm, mst, query)


def _d(value: Any) -> Decimal:
    return Decimal(str(value))


def test_summary(w: World) -> None:
    api, s = w.s.api, w.s
    summary = api.get(f"/reports/summary?{w.query}", token=s.su, salon=s.id)
    assert _d(summary["revenue"]["value"]) == 1500 and _d(summary["avgCheck"]["value"]) == 750
    assert (
        _d(summary["revenue"]["previous"]) == 300 and summary["revenue"]["changePercent"] == 400.0
    )
    assert _d(summary["records"]["value"]) == 4 and _d(summary["clients"]["value"]) == 4
    admin_view = api.get(f"/reports/summary?{w.query}", token=w.adm, salon=s.id)
    assert admin_view["revenue"] is None and admin_view["avgCheck"] is None
    api.get(f"/reports/summary?{w.query}", token=w.mst, salon=s.id, expect=403)


def test_revenue(w: World) -> None:
    api, s = w.s.api, w.s
    by_day = api.get(f"/reports/revenue?{w.query}&groupBy=day", token=s.su, salon=s.id)
    assert len(by_day["points"]) == 7 and _d(by_day["total"]) == 1500
    by_month = api.get(f"/reports/revenue?{w.query}&groupBy=month", token=s.su, salon=s.id)
    assert sum(_d(p["amount"]) for p in by_month["points"]) == 1500
    api.get(f"/reports/revenue?{w.query}", token=w.adm, salon=s.id, expect=403)


def test_clients(w: World) -> None:
    clients = w.s.api.get(f"/reports/clients?{w.query}&groupBy=week", token=w.adm, salon=w.s.id)
    assert (clients["total"], clients["new"], clients["returning"]) == (3, 2, 1)


def test_staff(w: World) -> None:
    api, s = w.s.api, w.s
    rows = {r["name"]: r for r in api.get(f"/reports/staff?{w.query}", token=s.su, salon=s.id)}
    anna, bohdana = rows["Анна"], rows["Богдана"]
    assert _d(anna["revenue"]) == 1200 and anna["completed"] == 2 and anna["rating"] == 5.0
    assert bohdana["noShow"] == 1 and _d(bohdana["revenue"]) == 300
    assert all(
        r["revenue"] is None for r in api.get(f"/reports/staff?{w.query}", token=w.adm, salon=s.id)
    )


def test_services(w: World) -> None:
    data = w.s.api.get(f"/reports/services?{w.query}", token=w.s.su, salon=w.s.id)
    by_name = {x["name"]: x for x in data["services"]}
    # Чек 1200 за «Стрижка + Фарбування» делится пропорционально цене: 300 / 900
    assert _d(by_name["Стрижка"]["revenue"]) == 600 and by_name["Стрижка"]["count"] == 2
    assert _d(by_name["Фарбування"]["revenue"]) == 900 and by_name["Фарбування"]["count"] == 2
    assert {c["category"]["name"] for c in data["categories"]} == {"Волосся", "Фарбування"}
    assert by_name["Стрижка"]["category"]["name"] == "Волосся"


def test_period_validation_and_export(w: World) -> None:
    api, s = w.s.api, w.s
    api.get(
        "/reports/staff?dateFrom=2026-02-01T00:00:00&dateTo=2026-01-01T00:00:00",
        token=s.su,
        salon=s.id,
        expect=422,
    )
    admin_file = api.get(f"/reports/export?{w.query}", token=w.adm, salon=s.id)
    assert load_workbook(io.BytesIO(admin_file.content)).sheetnames == [
        "Підсумки",
        "Клієнти",
        "Співробітники",
        "Послуги",
    ]
    su_file = api.get(f"/reports/export?{w.query}", token=s.su, salon=s.id)
    assert "Виручка" in load_workbook(io.BytesIO(su_file.content)).sheetnames
