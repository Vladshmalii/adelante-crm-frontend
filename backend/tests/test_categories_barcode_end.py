"""Категории услуг, штрихкод товара, сортировка по марже, ручной конец записи."""

import io
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from typing import Any
from zoneinfo import ZoneInfo

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import BOT_KEY, Salon, phone, setup_shifts

TZ = ZoneInfo("Europe/Kyiv")
BOT = {"prefix": "/api/bot", "headers": {"X-API-Key": BOT_KEY}}


@dataclass
class World:
    s: Salon
    adm: str
    master: dict[str, Any]
    mst: str
    other_mst: str


@pytest.fixture(scope="module")
def w(new_salon: Salon) -> World:
    s = new_salon
    _, adm = s.create_staff("administrator")
    master, mst = s.create_staff(firstName="Анна")
    _, other_mst = s.create_staff(firstName="Богдана")
    return World(s, adm, master, mst, other_mst)


# --- 1. Категории услуг ---------------------------------------------------------------


def test_service_categories(w: World) -> None:
    api, s = w.s.api, w.s
    initial = api.get("/services/categories", token=w.mst, salon=s.id)  # мастер читает
    assert [(c["name"], c["isSystem"]) for c in initial] == [("Інше", True)]
    other_id = initial[0]["id"]

    for name in ("Масаж", "Брови", "Ґель-лак", "Інтимна епіляція"):
        api.post("/services/categories", token=w.adm, salon=s.id, expect=201, json={"name": name})
    api.post("/services/categories", token=w.adm, salon=s.id, expect=409, json={"name": "масаж"})
    api.post("/services/categories", token=w.mst, salon=s.id, expect=403, json={"name": "X"})
    names = [c["name"] for c in api.get("/services/categories", token=w.adm, salon=s.id)]
    # Украинский алфавит (Ґ после Г, І после И), «Інше» — последней
    assert names == ["Брови", "Ґель-лак", "Інтимна епіляція", "Масаж", "Інше"]
    cats = {c["name"]: c["id"] for c in api.get("/services/categories", token=w.adm, salon=s.id)}

    renamed = api.patch(
        f"/services/categories/{cats['Брови']}",
        token=w.adm,
        salon=s.id,
        json={"name": "Брови та вії"},
    )
    assert renamed["name"] == "Брови та вії"
    api.patch(
        f"/services/categories/{cats['Масаж']}",
        token=w.adm,
        salon=s.id,
        expect=409,
        json={"name": "брови та вії"},
    )
    api.patch(
        f"/services/categories/{other_id}", token=w.adm, salon=s.id, expect=409, json={"name": "X"}
    )
    api.delete(f"/services/categories/{other_id}", token=w.adm, salon=s.id, expect=409)

    plain = api.post(
        "/services", token=w.adm, salon=s.id, expect=201, json={"name": "Консультація", "price": 0}
    )
    assert plain["category"] == {"id": other_id, "name": "Інше"}
    massage = api.post(
        "/services",
        token=w.adm,
        salon=s.id,
        expect=201,
        json={"name": "Масаж спини", "price": 500, "categoryId": cats["Масаж"]},
    )
    old = api.post(
        "/services",
        token=w.adm,
        salon=s.id,
        expect=201,
        json={"name": "Старий масаж", "price": 400, "categoryId": cats["Масаж"]},
    )
    api.delete(f"/services/{old['id']}", token=w.adm, salon=s.id)  # в архив
    api.post(
        "/services",
        token=w.adm,
        salon=s.id,
        expect=422,
        json={"name": "X", "price": 1, "categoryId": "00000000-0000-0000-0000-000000000000"},
    )
    counts = {
        c["name"]: c["servicesCount"]
        for c in api.get("/services/categories", token=w.adm, salon=s.id)
    }
    assert counts["Масаж"] == 1 and counts["Інше"] == 1  # архивные не считаются
    filtered = api.get(f"/services?categoryId={cats['Масаж']}", token=w.adm, salon=s.id)
    assert [x["name"] for x in filtered] == ["Масаж спини"]
    moved = api.patch(
        f"/services/{plain['id']}", token=w.adm, salon=s.id, json={"categoryId": cats["Масаж"]}
    )
    assert moved["category"]["name"] == "Масаж"

    # Удаление категории: все услуги (и архивные) — в «Інше»
    api.delete(f"/services/categories/{cats['Масаж']}", token=w.adm, salon=s.id, expect=204)
    archived = api.get(f"/services?categoryId={other_id}&status=archived", token=w.adm, salon=s.id)
    assert [x["name"] for x in archived] == ["Старий масаж"]
    active = api.get(f"/services?categoryId={other_id}", token=w.adm, salon=s.id)
    assert {x["name"] for x in active} == {"Консультація", "Масаж спини"}
    assert massage["id"] in {x["id"] for x in active}
    audit = api.get("/audit?entity=service", token=w.adm, salon=s.id)
    assert any(
        a["entityName"] == "Категорія послуг «Масаж»" and a["action"] == "deleted" for a in audit
    )


def test_booking_services_sorted_by_category(w: World) -> None:
    api, s = w.s.api, w.s
    cats = {c["name"]: c["id"] for c in api.get("/services/categories", token=w.adm, salon=s.id)}
    for name, category in (
        ("Ламінування", "Брови та вії"),
        ("Корекція", "Брови та вії"),
        ("Гель", "Ґель-лак"),
    ):
        api.post(
            "/services",
            token=w.adm,
            salon=s.id,
            expect=201,
            json={
                "name": name,
                "price": 100,
                "categoryId": cats[category],
                "masterIds": [w.master["id"]],
            },
        )
    api.post(
        "/services",
        token=w.adm,
        salon=s.id,
        expect=201,
        json={"name": "Діагностика", "price": 100, "masterIds": [w.master["id"]]},
    )
    listed = api.get("/services", prefix=f"/api/booking/{s.slug}")
    assert [(x["category_name"], x["name"]) for x in listed] == [
        ("Брови та вії", "Корекція"),
        ("Брови та вії", "Ламінування"),
        ("Ґель-лак", "Гель"),
        ("Інше", "Діагностика"),
    ]
    assert listed[0]["category_id"] == cats["Брови та вії"] and "category" not in listed[0]


# --- 2. Штрихкод ---------------------------------------------------------------------------


def _product(w: World, expect: int = 201, **body: Any) -> Any:
    return w.s.api.post(
        "/inventory/products",
        token=w.adm,
        salon=w.s.id,
        expect=expect,
        json={"unit": "pcs", **body},
    )


def test_barcode(w: World) -> None:
    api, s = w.s.api, w.s
    a = _product(w, name="Шампунь", sku="BC-1", barcode="  4820000123456  ")
    assert a["barcode"] == "4820000123456"
    assert _product(w, name="Без коду", sku="BC-2", barcode="   ")["barcode"] is None
    _product(w, expect=409, name="Дубль", sku="BC-3", barcode="4820000123456")
    _product(w, expect=422, name="Довгий", sku="BC-4", barcode="1" * 65)
    letters = _product(w, name="Маска", sku="BC-5", barcode="ABC-x")
    api.delete(f"/inventory/products/{letters['id']}", token=w.adm, salon=s.id, expect=204)
    _product(w, expect=409, name="Регістр", sku="BC-6", barcode="abc-X")  # неактивный тоже

    found = api.get("/inventory/products?query=3456", token=w.adm, salon=s.id)
    assert [p["name"] for p in found] == ["Шампунь"]
    for_master = api.get("/inventory/products?query=3456", token=w.mst, salon=s.id)
    assert [p["name"] for p in for_master] == ["Шампунь"] and for_master[0]["barcode"] is None

    api.patch(
        f"/inventory/products/{a['id']}",
        token=w.adm,
        salon=s.id,
        expect=409,
        json={"barcode": "ABC-X"},
    )
    cleared = api.patch(
        f"/inventory/products/{a['id']}", token=w.adm, salon=s.id, json={"barcode": None}
    )
    assert cleared["barcode"] is None
    api.patch(
        f"/inventory/products/{a['id']}", token=w.adm, salon=s.id, json={"barcode": "4820000123456"}
    )


def _xlsx(rows: list[list[Any]]) -> bytes:
    wb = Workbook()
    for row in rows:
        wb.active.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _import(w: World, rows: list[list[Any]]) -> Any:
    return w.s.api.post(
        "/inventory/import",
        token=w.adm,
        salon=w.s.id,
        files={
            "file": (
                "p.xlsx",
                _xlsx(rows),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )


def test_barcode_import_export(w: World) -> None:
    api, s = w.s.api, w.s
    report = _import(
        w,
        [
            ["Назва", "Артикул", "Штрихкод", "Категорія", "Одиниця", "Залишок"],
            ["Шампунь", "BC-1", "", None, "шт", 5],  # пустая ячейка — штрихкод не стирается
            ["Кондиціонер", "IM-1", " 111222333 ", None, "шт", 3],
            ["Дубль", "IM-2", "4820000123456", None, "шт", 1],  # штрихкод у BC-1
            ["Дубль у файлі", "IM-3", "111222333", None, "шт", 1],  # только что у IM-1
        ],
    )
    assert report["created"] == 1 and report["updated"] == 1
    assert len(report["errors"]) == 2 and "Шампунь" in report["errors"][0]
    rows = {
        p["sku"]: p
        for p in api.get(
            "/inventory/products?perPage=500&includeInactive=true", token=w.adm, salon=s.id
        )
    }
    assert rows["BC-1"]["barcode"] == "4820000123456"
    assert rows["IM-1"]["barcode"] == "111222333"
    assert "IM-2" not in rows and "IM-3" not in rows

    # Старый шаблон без колонки «Штрихкод» — по заголовкам
    legacy = _import(
        w,
        [
            ["Назва", "Артикул", "Категорія", "Одиниця", "Об'єм упаковки", "Залишок"],
            ["Старий шаблон", "LG-1", None, "мл", 500, 100],
        ],
    )
    assert legacy["created"] == 1 and legacy["errors"] == []
    # Порядок колонок любой
    shuffled = _import(w, [["Штрихкод", "Артикул", "Назва"], ["777", "SH-9", "Перемішаний"]])
    assert shuffled["created"] == 1
    assert api.get("/inventory/products?query=777", token=w.adm, salon=s.id)[0]["sku"] == "SH-9"

    export = api.get("/inventory/export?blocks=main", token=w.adm, salon=s.id)
    sheet = load_workbook(io.BytesIO(export.content)).active
    header = [c.value for c in sheet[1]]
    assert header == ["Назва", "Артикул", "Штрихкод", "Категорія", "Одиниця"]
    exported = {row[1]: row[2] for row in sheet.iter_rows(min_row=2, values_only=True)}
    assert exported["IM-1"] == "111222333"


# --- 3. Сортировка по марже ----------------------------------------------------------------


def test_sort_by_margin(w: World) -> None:
    api, s = w.s.api, w.s
    category = api.post(
        "/inventory/categories", token=w.adm, salon=s.id, expect=201, json={"name": "Маржа"}
    )["id"]
    for name, cost, price in (
        ("Б-середня", 50, 100),  # 0.5
        ("А-висока", 10, 100),  # 0.9
        ("В-низька", 90, 100),  # 0.1
        ("Г-без ціни", 10, None),
        ("Д-нульова ціна", 10, 0),
        ("Е-друга середня", 25, 50),  # 0.5 — вторично по названию
    ):
        _product(
            w, name=name, sku=f"MG-{name}", costPrice=cost, salePrice=price, categoryId=category
        )

    def order(desc: bool) -> list[str]:
        rows = api.get(
            f"/inventory/products?categoryId={category}&sort=margin&desc={desc}",
            token=w.adm,
            salon=s.id,
        )
        return [p["name"] for p in rows]

    assert order(False) == [
        "В-низька",
        "Б-середня",
        "Е-друга середня",
        "А-висока",
        "Г-без ціни",
        "Д-нульова ціна",
    ]
    assert order(True) == [
        "А-висока",
        "Б-середня",
        "Е-друга середня",
        "В-низька",
        "Г-без ціни",
        "Д-нульова ціна",
    ]


# --- 4. Конец записи -----------------------------------------------------------------------------


def _day() -> date:
    return datetime.now(TZ).date() + timedelta(days=3)


def _at(day: date, hour: int, minute: int = 0) -> str:
    return datetime(day.year, day.month, day.day, hour, minute).isoformat()


@pytest.fixture(scope="module")
def svc(w: World) -> dict[str, Any]:
    setup_shifts(
        w.s,
        [w.master["id"]],
        datetime.now(TZ).date(),
        _day() + timedelta(days=5),
        start="09:00",
        end="18:00",
        break_start=None,
        break_end=None,
    )
    return w.s.api.post(
        "/services",
        token=w.adm,
        salon=w.s.id,
        expect=201,
        json={"name": "Фарбування", "price": 900, "durationMinutes": 60},
    )


def _record(w: World, svc: dict[str, Any], expect: int = 201, **body: Any) -> Any:
    return w.s.api.post(
        "/records",
        token=w.adm,
        salon=w.s.id,
        expect=expect,
        json={
            "newClient": {"name": "К", "phone": phone()},
            "masterId": w.master["id"],
            "serviceIds": [svc["id"]],
            **body,
        },
    )


def _local(value: str) -> str:
    return datetime.fromisoformat(value).astimezone(TZ).strftime("%d %H:%M")


def test_create_with_end(w: World, svc: dict[str, Any]) -> None:
    day = _day()
    rec = _record(w, svc, startAt=_at(day, 10), endAt=_at(day, 12, 30))
    assert _local(rec["endAt"]) == f"{day:%d} 12:30" and float(rec["totalAmount"]) == 900
    default = _record(w, svc, startAt=_at(day, 10))
    assert _local(default["endAt"]) == f"{day:%d} 11:00"
    _record(w, svc, expect=422, startAt=_at(day, 10), endAt=_at(day, 10))
    _record(w, svc, expect=422, startAt=_at(day, 23), endAt=_at(day + timedelta(days=1), 0, 30))
    midnight = _record(w, svc, startAt=_at(day, 23), endAt=_at(day + timedelta(days=1), 0))
    assert _local(midnight["endAt"]) == f"{day + timedelta(days=1):%d} 00:00"


def test_patch_end(w: World, svc: dict[str, Any], monkeypatch: pytest.MonkeyPatch) -> None:
    from workers import telegram
    from workers.celery_app import celery
    from workers.tasks import outbox

    api, s = w.s.api, w.s
    sent: list[tuple[int, str]] = []
    monkeypatch.setattr(celery.conf, "task_always_eager", True)
    monkeypatch.setattr(telegram, "send_message", lambda chat, text: sent.append((chat, text)))
    master_phone = api.get("/auth/me", token=w.mst)["phone"]
    chat = 980_000_000 + int("".join(ch for ch in master_phone if ch.isdigit())[-5:])
    api.post("/link-telegram", **BOT, json={"phone": master_phone, "telegram_user_id": chat})
    day = _day() + timedelta(days=1)
    rec = _record(w, svc, startAt=_at(day, 10))
    outbox.publish_outbox()
    sent.clear()

    # Только конец — растягивание карточки: аудит, без уведомления мастеру
    stretched = api.patch(
        f"/records/{rec['id']}", token=w.adm, salon=s.id, json={"endAt": _at(day, 11, 45)}
    )
    assert _local(stretched["endAt"]) == f"{day:%d} 11:45"
    assert float(stretched["totalAmount"]) == 900 and stretched["outsideShift"] is False
    outbox.publish_outbox()
    assert [c for c, _ in sent if c == chat] == []
    history = api.get(f"/records/{rec['id']}", token=w.adm, salon=s.id)["history"]
    assert any("endAt" in (h["details"] or {}) for h in history)

    # Перенос только startAt — ручная длительность сохраняется; мастеру уведомление
    moved = api.patch(
        f"/records/{rec['id']}", token=w.adm, salon=s.id, json={"startAt": _at(day, 13)}
    )
    assert _local(moved["endAt"]) == f"{day:%d} 14:45"
    outbox.publish_outbox()
    assert any("Запис змінено" in text for c, text in sent if c == chat)

    # Растянуть за верхний край: startAt + endAt; за пределы смены — outsideShift
    both = api.patch(
        f"/records/{rec['id']}",
        token=w.mst,
        salon=s.id,
        json={"startAt": _at(day, 16), "endAt": _at(day, 19)},
    )
    assert _local(both["startAt"]) == f"{day:%d} 16:00" and both["outsideShift"] is True

    # Тот же состав услуг (форма шлёт его всегда) — ручная длительность не сбрасывается
    same = api.patch(
        f"/records/{rec['id']}",
        token=w.adm,
        salon=s.id,
        json={"serviceIds": [svc["id"]], "startAt": _at(day, 9)},
    )
    assert _local(same["endAt"]) == f"{day:%d} 12:00"
    # Другой состав услуг без endAt — конец снова по услугам
    two = api.post(
        "/services",
        token=w.adm,
        salon=s.id,
        expect=201,
        json={"name": "Укладка", "price": 300, "durationMinutes": 30},
    )
    reset = api.patch(
        f"/records/{rec['id']}",
        token=w.adm,
        salon=s.id,
        json={"serviceIds": [svc["id"], two["id"]]},
    )
    assert _local(reset["endAt"]) == f"{day:%d} 10:30"
    with_end = api.patch(
        f"/records/{rec['id']}",
        token=w.adm,
        salon=s.id,
        json={"serviceIds": [svc["id"]], "endAt": _at(day, 12)},
    )
    assert _local(with_end["endAt"]) == f"{day:%d} 12:00"
    assert float(with_end["totalAmount"]) == 900  # цена — по услугам, не по длительности

    api.patch(
        f"/records/{rec['id']}", token=w.adm, salon=s.id, expect=422, json={"endAt": _at(day, 8)}
    )
    api.patch(
        f"/records/{rec['id']}",
        token=w.adm,
        salon=s.id,
        expect=422,
        json={"endAt": _at(day + timedelta(days=1), 1)},
    )
    api.patch(
        f"/records/{rec['id']}",
        token=w.other_mst,
        salon=s.id,
        expect=404,
        json={"endAt": _at(day, 11)},
    )
    api.post(f"/records/{rec['id']}/status", token=w.adm, salon=s.id, json={"status": "cancelled"})
    api.patch(
        f"/records/{rec['id']}", token=w.adm, salon=s.id, expect=409, json={"endAt": _at(day, 11)}
    )
