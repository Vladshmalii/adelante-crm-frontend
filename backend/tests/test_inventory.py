"""Склад: товары, движения, категории, импорт/экспорт, расходники записи."""

import io
from dataclasses import dataclass
from decimal import Decimal
from typing import Any

import pytest
from openpyxl import Workbook, load_workbook

from tests.conftest import Salon, phone


@dataclass
class World:
    s: Salon
    adm: str
    master: dict[str, Any]
    mst: str
    master2: dict[str, Any]
    mst2: str
    system: dict[str, Any]
    categories: list[dict[str, Any]]

    def product(self, **body: Any) -> dict[str, Any]:
        return self.s.api.post(
            "/inventory/products", token=self.adm, salon=self.s.id, expect=201, json=body
        )


@pytest.fixture(scope="module")
def w(new_salon: Salon) -> World:
    s = new_salon
    _, adm = s.create_staff("administrator")
    master, mst = s.create_staff()
    master2, mst2 = s.create_staff()
    categories = s.api.get("/inventory/categories", token=adm, salon=s.id)
    system = next(c for c in categories if c["isSystem"])
    return World(s, adm, master, mst, master2, mst2, system, categories)


def _qty(value: Any) -> Decimal:
    return Decimal(str(value))


def test_base_categories(w: World) -> None:
    assert [c["name"] for c in w.categories] == [
        "Без категорії",
        "Витратні матеріали",
        "Обладнання",
        "Професійна косметика",
        "Товари для продажу",
    ]
    w.s.api.get("/inventory/categories", token=w.mst, salon=w.s.id, expect=403)


def test_products_and_access(w: World) -> None:
    api, s = w.s.api, w.s
    shampoo = w.product(
        name="Шампунь",
        sku="SH-1",
        unit="ml",
        packageVolume=500,
        quantity=1500,
        minQuantity=600,
        costPrice=0.4,
        salePrice=0.9,
    )
    assert shampoo["costPrice"] is not None  # цены видит администратор
    assert shampoo["category"]["id"] == w.system["id"] and shampoo["stockStatus"] == "in_stock"
    gloves = w.product(
        name="Рукавички",
        sku="GL-1",
        quantity=10,
        minQuantity=20,
        costPrice=5,
        categoryId=w.categories[1]["id"],
    )
    assert gloves["stockStatus"] == "low"
    api.post(
        "/inventory/products",
        token=w.adm,
        salon=s.id,
        expect=409,
        json={"name": "Дубль", "sku": "sh-1"},
    )
    api.post(
        "/inventory/products", token=w.mst, salon=s.id, expect=403, json={"name": "X", "sku": "X"}
    )

    limited = api.get("/inventory/products", token=w.mst, salon=s.id)
    assert {p["name"] for p in limited} >= {"Шампунь", "Рукавички"}
    assert all(p["sku"] is None and p["costPrice"] is None for p in limited)
    api.get(f"/inventory/products/{shampoo['id']}", token=w.mst, salon=s.id, expect=403)
    assert [
        p["name"] for p in api.get("/inventory/products?stockStatus=low", token=w.adm, salon=s.id)
    ] == ["Рукавички"]
    assert [
        p["name"] for p in api.get("/inventory/products?query=gl-1", token=w.adm, salon=s.id)
    ] == ["Рукавички"]
    summary = api.get("/inventory/summary", token=w.adm, salon=s.id)
    assert summary["total"] == 2 and summary["low"] == 1
    assert _qty(summary["stockValue"]) == Decimal("650.00")


def _find(w: World, sku: str) -> dict[str, Any]:
    return w.s.api.get(f"/inventory/products?query={sku}", token=w.adm, salon=w.s.id)[0]


def test_movements(w: World) -> None:
    api, s = w.s.api, w.s
    shampoo = _find(w, "SH-1")
    base = f"/inventory/products/{shampoo['id']}/movements"
    out = api.post(
        base,
        token=w.adm,
        salon=s.id,
        expect=201,
        json={"type": "write_off", "quantity": 300, "reason": "Брак"},
    )
    assert _qty(out["delta"]) == -300 and _qty(out["quantityAfter"]) == 1200
    api.post(
        base, token=w.adm, salon=s.id, expect=409, json={"type": "write_off", "quantity": 5000}
    )
    adj = api.post(
        base,
        token=w.adm,
        salon=s.id,
        expect=201,
        json={"type": "adjustment", "quantity": 1000, "reason": "Інвентаризація"},
    )
    assert _qty(adj["delta"]) == -200 and _qty(adj["quantityAfter"]) == 1000
    api.post(base, token=w.adm, salon=s.id, expect=422, json={"type": "receipt", "quantity": 0})
    api.post(base, token=w.mst, salon=s.id, expect=403, json={"type": "receipt", "quantity": 1})
    history = api.get(base, token=w.adm, salon=s.id)
    assert [h["type"] for h in history] == ["adjustment", "write_off", "receipt"]

    patched = api.patch(
        f"/inventory/products/{shampoo['id']}",
        token=w.adm,
        salon=s.id,
        json={"salePrice": 1.1, "quantity": 99999},
    )
    assert _qty(patched["quantity"]) == 1000 and _qty(patched["salePrice"]) == Decimal("1.1")
    api.patch(
        f"/inventory/products/{shampoo['id']}",
        token=w.adm,
        salon=s.id,
        expect=409,
        json={"sku": "GL-1"},
    )


def test_categories(w: World) -> None:
    api, s = w.s.api, w.s
    cat = api.post(
        "/inventory/categories", token=w.adm, salon=s.id, expect=201, json={"name": "Фарби"}
    )
    api.post("/inventory/categories", token=w.adm, salon=s.id, expect=409, json={"name": "фарби"})
    api.patch(
        f"/inventory/categories/{cat['id']}",
        token=w.adm,
        salon=s.id,
        json={"name": "Фарби для волосся"},
    )
    api.patch(
        f"/inventory/categories/{w.system['id']}",
        token=w.adm,
        salon=s.id,
        expect=409,
        json={"name": "X"},
    )
    api.delete(f"/inventory/categories/{w.system['id']}", token=w.adm, salon=s.id, expect=409)

    dye = w.product(name="Фарба", sku="DY-1", categoryId=cat["id"], quantity=3)
    api.delete(
        f"/inventory/categories/{cat['id']}?mode=move_to_uncategorized",
        token=w.adm,
        salon=s.id,
        expect=204,
    )
    moved = api.get(f"/inventory/products/{dye['id']}", token=w.adm, salon=s.id)
    assert moved["category"]["id"] == w.system["id"] and moved["isActive"]

    tmp_cat = api.post(
        "/inventory/categories", token=w.adm, salon=s.id, expect=201, json={"name": "Тимчасова"}
    )
    tmp = w.product(name="Тимчасовий", sku="TM-1", categoryId=tmp_cat["id"])
    api.delete(
        f"/inventory/categories/{tmp_cat['id']}?mode=delete_products",
        token=w.adm,
        salon=s.id,
        expect=204,
    )
    assert api.get(f"/inventory/products/{tmp['id']}", token=w.adm, salon=s.id)["isActive"] is False


def test_import_export(w: World) -> None:
    api, s = w.s.api, w.s
    wb = Workbook()
    ws = wb.active
    ws.append(
        [
            "Назва",
            "Артикул",
            "Категорія",
            "Одиниця",
            "Об'єм упаковки",
            "Залишок",
            "Мін. залишок",
            "Собівартість",
            "Ціна продажу",
            "Опис",
        ]
    )
    ws.append(
        [
            "Шампунь (імпорт)",
            "SH-1",
            "Професійна косметика",
            "мл",
            500,
            1250,
            500,
            0.45,
            1.2,
            "Оновлено",
        ]
    )
    ws.append(["Маска", "MK-1", "Нова категорія", "г", 250, 750, 100, 0.3, None, None])
    ws.append(["", "XX-1", None, None, None, None, None, None, None, None])
    ws.append(["Кривий", "KR-1", None, "бочка", None, None, None, None, None, None])
    ws.append(["Кривий2", "KR-2", None, "шт", None, "багато", None, None, None, None])
    buffer = io.BytesIO()
    wb.save(buffer)
    report = api.post(
        "/inventory/import",
        token=w.adm,
        salon=s.id,
        files={
            "file": (
                "p.xlsx",
                buffer.getvalue(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
        },
    )
    assert report["created"] == 1 and report["updated"] == 1 and len(report["errors"]) == 3
    shampoo = _find(w, "SH-1")
    assert _qty(shampoo["quantity"]) == 1250
    assert shampoo["category"]["name"] == "Професійна косметика"
    names = [c["name"] for c in api.get("/inventory/categories", token=w.adm, salon=s.id)]
    assert "Нова категорія" in names

    export = api.get("/inventory/export?blocks=main,stock", token=w.adm, salon=s.id)
    header = [cell.value for cell in load_workbook(io.BytesIO(export.content)).active[1]]
    assert header == [
        "Назва",
        "Артикул",
        "Категорія",
        "Одиниця",
        "Об'єм упаковки",
        "Залишок",
        "Мін. залишок",
        "Статус",
    ]
    api.get("/inventory/export", token=w.mst, salon=s.id, expect=403)


def test_consumables(w: World) -> None:
    api, s = w.s.api, w.s
    svc = api.post(
        "/services",
        token=w.adm,
        salon=s.id,
        expect=201,
        json={"name": "Стрижка", "price": 300, "durationMinutes": 30},
    )
    shampoo, gloves = _find(w, "SH-1"), _find(w, "GL-1")

    def record(master: dict[str, Any]) -> dict[str, Any]:
        return api.post(
            "/records",
            token=w.adm,
            salon=s.id,
            expect=201,
            json={
                "newClient": {"name": "К", "phone": phone()},
                "masterId": master["id"],
                "serviceIds": [svc["id"]],
                "startAt": "2030-03-01T10:00:00",
            },
        )

    own, other = record(w.master), record(w.master2)
    items = api.post(
        f"/records/{own['id']}/consumables",
        token=w.mst,
        salon=s.id,
        expect=201,
        json={
            "items": [
                {"productId": shampoo["id"], "quantity": 50},
                {"productId": shampoo["id"], "quantity": 25},
                {"productId": gloves["id"], "quantity": 2},
            ]
        },
    )
    assert sorted(_qty(i["quantity"]) for i in items) == [Decimal(2), Decimal(75)]
    api.post(
        f"/records/{other['id']}/consumables",
        token=w.mst,
        salon=s.id,
        expect=404,
        json={"items": [{"productId": gloves["id"], "quantity": 1}]},
    )
    api.post(
        f"/records/{other['id']}/consumables",
        token=w.adm,
        salon=s.id,
        expect=409,
        json={"items": [{"productId": gloves["id"], "quantity": 1000}]},
    )
    api.post(
        f"/records/{other['id']}/consumables",
        token=w.adm,
        salon=s.id,
        expect=201,
        json={"items": [{"productId": gloves["id"], "quantity": 1}]},
    )
    assert len(api.get(f"/records/{other['id']}/consumables", token=w.mst2, salon=s.id)) == 1
    api.get(f"/records/{other['id']}/consumables", token=w.mst, salon=s.id, expect=404)
    assert _qty(_find(w, "GL-1")["quantity"]) == 7
    history = api.get(f"/inventory/products/{gloves['id']}/movements", token=w.adm, salon=s.id)
    assert history[0]["recordId"] == other["id"]

    api.post(
        f"/records/{other['id']}/status", token=w.adm, salon=s.id, json={"status": "cancelled"}
    )
    api.post(
        f"/records/{other['id']}/consumables",
        token=w.adm,
        salon=s.id,
        expect=409,
        json={"items": [{"productId": gloves["id"], "quantity": 1}]},
    )
